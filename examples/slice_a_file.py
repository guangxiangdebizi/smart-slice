"""Slice real files: a generated .docx from disk, and bytes with image capture.

Builds a small .docx in a temp dir (needs the `office` extra), slices it via
slice_path, then demonstrates slice_bytes + the save_image callback on a document
whose paragraphs reference extracted images.

Run: python examples/slice_a_file.py
"""
import os
import tempfile

from smart_slice import ImageAsset, slice_bytes, slice_path


def _make_docx(path: str) -> None:
    """Write a minimal two-heading .docx (python-docx)."""
    from docx import Document

    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew steadily across all regions this quarter.")
    doc.add_heading("Highlights", level=2)
    doc.add_paragraph("Cloud subscriptions expanded by a double-digit margin.")
    doc.save(path)


def _markdown_with_image() -> bytes:
    """Markdown that already carries an ./oss/file/{id} image reference."""
    image_id = "11111111-1111-7111-8111-111111111111"
    text = (
        "# Field Notes\n\n"
        "See the diagram below.\n\n"
        f"![diagram](./oss/file/{image_id})\n\n"
        "The diagram explains the flow.\n"
    )
    return text.encode("utf-8")


def main() -> None:
    # 1) slice a real .docx from disk
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "report.docx")
        try:
            _make_docx(docx_path)
        except ImportError:
            print("python-docx not installed; run: pip install smart-slice[office]")
            return
        rows = slice_path(docx_path, limit=1000)
        print(f"== slice_path('report.docx') -> {len(rows)} paragraphs ==")
        for row in rows:
            print(f"  [{row['title'] or '(no heading)'}] {row['content'][:50]!r}")

    print()

    # 2) slice bytes and capture images via the save_image callback
    captured = []

    def save_image(assets):
        # assets: list[ImageAsset]; return {new_id: existing_id} to remap refs
        captured.extend(assets)
        return None

    data = _markdown_with_image()
    rows = slice_bytes(data, "notes.md", limit=1000, save_image=save_image)
    print(f"== slice_bytes('notes.md') -> {len(rows)} paragraphs, {len(captured)} images captured ==")
    for row in rows:
        print(f"  [{row['title'] or '(no heading)'}] {row['content'][:50]!r}")
    # a markdown file has no embedded binary images, so nothing is captured here;
    # a .docx/.pptx/.pdf with real images would populate `captured` with ImageAsset
    # objects whose .content holds the raw bytes.
    if captured:
        for asset in captured:
            assert isinstance(asset, ImageAsset)
            print(f"  captured image id={asset.id} bytes={len(asset.content)}")


if __name__ == "__main__":
    main()