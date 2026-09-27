"""PptxSplitHandle（.pptx）：python-pptx 逐页提取文本框/表格/演讲者备注。

Phase 2-A 全格式适配（2026-09-17）：每页（slide）组织为一个语义单元，页内以
标题占位符文字（无则"第N页"）作为小节标题，正文 = 文本框段落 + 表格（Markdown
管道表）+ 演讲者备注（"备注："前缀段落）。返回结构与 docx/html 等单文件 handler
一致：{"name", "content": [{title, content}, ...]}，文本统一交给 SplitModel
（标题树 + limit 二次切分）。

内嵌图片：本期与 docx 现状对齐，图片占位 markdown 引用经 save_image 回调收集落库，
不做内嵌图 OCR（下一分支 feat/inline-image-ocr 接入）。
"""
# coding=utf-8
# Derived from the source platform document-slicing layer (see docs/PORTING.md).
# Imports and framework-facing symbols were re-routed; the parsing and
# splitting logic is unchanged.
from smart_slice._logging import get_logger

_log = get_logger("pptx")


from smart_slice._i18n import gettext as _
from smart_slice.exceptions import SliceError, ResourceLimitError
from smart_slice.types import ImageAsset
from smart_slice import _uuid as uuid

import io
import os
import traceback
from typing import List

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:  # optional extra: smart-slice[office]
    Presentation = None
    PPTX_AVAILABLE = False

from smart_slice.handlers.base import BaseSplitHandle
from smart_slice._validation import validate_ooxml
from smart_slice.handlers._utils import build_split_model, md_table


PRESENTATION_CONTENT_TYPES = {
    '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',
    '.pptm': 'application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml',
    '.ppsx': 'application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml',
    '.ppsm': 'application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml',
    '.potx': 'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml',
    '.potm': 'application/vnd.ms-powerpoint.template.macroEnabled.main+xml',
}


def _text_frame_text(text_frame) -> str:
    lines = ["".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
             for paragraph in text_frame.paragraphs]
    return "\n".join(line for line in lines if line and line.strip())


def _shape_image_md(shape, image_list) -> str:
    try:
        image = shape.image
        image_id = str(uuid.uuid7())
        file_name = image.filename or f"image_{image_id}.png"
        image_list.append(ImageAsset(id=image_id, file_name=file_name,
                               meta={'debug': False, 'content': image.blob}))
        return f'![{file_name}](./oss/file/{image_id})'
    except Exception as e:
        _log.warning(f'pptx image extract failed: {e}')
        return ""


def _shape_text(shape, image_list) -> str:
    """提取单个 shape 的文本/表格/图片占位，返回 markdown 片段。"""
    if shape.has_text_frame:
        return _text_frame_text(shape.text_frame)
    if getattr(shape, 'has_table', False):
        rows = []
        for row in shape.table.rows:
            rows.append([cell.text.replace("\n", '</br>') for cell in row.cells])
        return md_table(rows)
    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
        return _shape_image_md(shape, image_list)
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP：递归子 shape
        return "\n\n".join(
            piece for piece in (_shape_text(sub, image_list) for sub in shape.shapes)
            if piece and piece.strip()
        )
    return ""


def _slide_section(slide, index, image_list) -> str:
    """单页 → markdown 小节：标题占位符为节标题，正文=文本框/表格/备注。"""
    pieces = []
    title_text = ""
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.has_text_frame:
        title_text = _text_frame_text(title_shape.text_frame).strip().split("\n")[0]
    if not title_text:
        title_text = f"第{index}页"
    for shape in slide.shapes:
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue  # 标题占位符已作节标题，不重复进正文
        text = _shape_text(shape, image_list)
        if text and text.strip():
            pieces.append(text.strip())
    # 演讲者备注
    if slide.has_notes_slide:
        notes_text = _text_frame_text(slide.notes_slide.notes_text_frame)
        if notes_text and notes_text.strip():
            pieces.append(f"备注：{notes_text.strip()}")
    body = "\n\n".join(pieces)
    return f"# {title_text}\n\n{body}" if body else f"# {title_text}"


def _presentation_to_md(buffer: bytes, image_list, name='') -> str:
    content_type = PRESENTATION_CONTENT_TYPES.get(os.path.splitext(str(name))[1].lower())
    validate_ooxml(buffer, 'ppt/presentation.xml',
                   (content_type,) if content_type else tuple(PRESENTATION_CONTENT_TYPES.values()))
    presentation = Presentation(io.BytesIO(buffer))
    return "\n\n".join(
        _slide_section(slide, index, image_list)
        for index, slide in enumerate(presentation.slides, start=1)
    )


class PptxSplitHandle(BaseSplitHandle):
    def support(self, file, get_buffer):
        # python-pptx 为可选 extra（smart-slice[office]）：缺失时不认领
        return PPTX_AVAILABLE and file.name.lower().endswith(tuple(PRESENTATION_CONTENT_TYPES))

    def handle(self, file, pattern_list: List, with_filter: bool, limit: int, get_buffer, save_image):
        image_list = []
        try:
            buffer = get_buffer(file)
            content = _presentation_to_md(buffer, image_list, file.name)
            if len(image_list) > 0:
                save_image(image_list)
            split_model = build_split_model(pattern_list, with_filter, limit)
        except (SliceError, ResourceLimitError):
            raise
        except BaseException as e:
            _log.error(f"Error processing PPTX file {file.name}: {e}, {traceback.format_exc()}")
            raise SliceError(
                500,
                _("Failed to parse document file {name}: {reason}").format(name=file.name, reason=e),
            )
        return {'name': file.name, 'content': split_model.parse(content)}

    def get_content(self, file, save_image):
        try:
            image_list = []
            content = _presentation_to_md(file.read(), image_list, getattr(file, 'name', ''))
            if len(image_list) > 0:
                save_image(image_list)
            return content
        except (SliceError, ResourceLimitError):
            raise
        except Exception as e:
            _log.error(f'Error getting pptx content: {e}', exc_info=True)
            raise SliceError(400, 'Invalid or damaged presentation') from e
