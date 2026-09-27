# coding=utf-8
# Derived from the source platform document-slicing layer (see docs/PORTING.md).
# Imports and framework-facing symbols were re-routed; the parsing and
# splitting logic is unchanged.
from smart_slice._logging import get_logger

_log = get_logger("xlsx")


from smart_slice._i18n import gettext as _
from smart_slice.exceptions import SliceError, ResourceLimitError

import io
import os
import traceback
from typing import List

import openpyxl
from openpyxl import load_workbook

from smart_slice.handlers.base import BaseSplitHandle
from smart_slice.handlers._xlsx_images import xlsx_embed_cells_images
from smart_slice._validation import ParserLimits, validate_ooxml
from smart_slice.chunker import smart_split_paragraph

splitter = '\n`-----------------------------------`\n'
WORKBOOK_CONTENT_TYPES = {
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
    '.xlsm': 'application/vnd.ms-excel.sheet.macroEnabled.main+xml',
    '.xltx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml',
    '.xltm': 'application/vnd.ms-excel.template.macroEnabled.main+xml',
}


def validate_workbook(buffer, name):
    content_type = WORKBOOK_CONTENT_TYPES.get(os.path.splitext(str(name))[1].lower())
    validate_ooxml(buffer, 'xl/workbook.xml',
                   (content_type,) if content_type else tuple(WORKBOOK_CONTENT_TYPES.values()))


def validate_worksheet_sizes(workbook):
    limits = ParserLimits.load()
    if sum(sheet.max_row * sheet.max_column for sheet in workbook.worksheets) > limits.max_table_cells:
        raise ResourceLimitError('Workbook exceeds the cell count limit')


def post_cell(image_dict, cell_value):
    image = image_dict.get(cell_value, None)
    if image is not None:
        return f'![](./oss/file/{image.id})'
    return cell_value.replace('\n', '<br>').replace('|', '&#124;')


def row_to_md(row, image_dict):
    return '| ' + ' | '.join(
        [post_cell(image_dict, str(cell.value if cell.value is not None else '')) if cell is not None else '' for cell
         in row]) + ' |\n'


def handle_sheet(file_name, sheet, image_dict, limit: int):
    rows = sheet.rows
    paragraphs = []
    result = {'name': file_name, 'content': paragraphs}
    try:
        title_row_list = next(rows)
        title_md_content = row_to_md(title_row_list, image_dict)
        title_md_content += '| ' + ' | '.join(
            ['---' if cell is not None else '' for cell in title_row_list]) + ' |\n'
    except Exception as e:
        return result
    if len(title_row_list) == 0:
        return result
    result_item_content = ''
    for row in rows:
        next_md_content = row_to_md(row, image_dict)
        next_md_content_len = len(next_md_content)
        result_item_content_len = len(result_item_content)
        if len(result_item_content) == 0:
            result_item_content += title_md_content
            result_item_content += next_md_content
        else:
            if result_item_content_len + next_md_content_len < limit:
                result_item_content += next_md_content
            else:
                paragraphs.append({'content': result_item_content, 'title': ''})
                result_item_content = title_md_content + next_md_content
        # 智能切片升级方案 1-4（2026-09-16，P7）：单行数据超 limit 时二次切分，行内容不丢失。
        # 累计路径已有 < limit 校验不会超限，此处超限必为「表头 + 单个超限行」形态
        if len(result_item_content) > limit:
            piece_limit = max(limit - len(title_md_content), 50)
            for piece in smart_split_paragraph(next_md_content, piece_limit):
                paragraphs.append({'content': title_md_content + piece, 'title': ''})
            result_item_content = ''
    if len(result_item_content) > 0:
        paragraphs.append({'content': result_item_content, 'title': ''})
    return result


class XlsxSplitHandle(BaseSplitHandle):
    def fill_merged_cells(self, sheet, image_dict):
        data = []

        # 获取第一行作为标题行
        headers = []
        for idx, cell in enumerate(sheet[1]):
            if cell.value is None:
                headers.append(' ' * (idx + 1))
            else:
                headers.append(cell.value)

        # 从第二行开始遍历每一行
        for row in sheet.iter_rows(min_row=2, values_only=False):
            row_data = {}
            for col_idx, cell in enumerate(row):
                cell_value = cell.value

                # 如果单元格为空，并且该单元格在合并单元格内，获取合并单元格的值
                if cell_value is None:
                    for merged_range in sheet.merged_cells.ranges:
                        if cell.coordinate in merged_range:
                            cell_value = sheet[merged_range.min_row][merged_range.min_col - 1].value
                            break

                image = image_dict.get(cell_value, None)
                if image is not None:
                    cell_value = f'![](./oss/file/{image.id})'

                # 使用标题作为键，单元格的值作为值存入字典
                row_data[headers[col_idx]] = cell_value
            data.append(row_data)

        return data

    def handle(self, file, pattern_list: List, with_filter: bool, limit: int, get_buffer, save_image):
        buffer = get_buffer(file)
        try:
            if type(limit) is str:
                limit = int(limit)
            validate_workbook(buffer, file.name)
            workbook = openpyxl.load_workbook(io.BytesIO(buffer), keep_links=False)
            validate_worksheet_sizes(workbook)
            # 矩阵验证修复 2026-09-16：公式单元格优先采用 Excel 保存时缓存的计算值。
            # 默认 data_only=False 读到公式串（如 '=SUM(B3:B4)'），切片后出现公式原文；
            # 用 data_only=True 工作簿的缓存值回填公式格，缓存缺失时保留公式串作为占位。
            try:
                workbook_data = openpyxl.load_workbook(io.BytesIO(buffer), data_only=True, keep_links=False)
                for ws_formula, ws_data in zip(workbook.worksheets, workbook_data.worksheets):
                    for row in ws_data.iter_rows():
                        for cell_data in row:
                            if cell_data.value is None:
                                continue
                            cell_formula = ws_formula[cell_data.coordinate]
                            if isinstance(cell_formula.value, str) and cell_formula.value.startswith('='):
                                cell_formula.value = cell_data.value
            except ResourceLimitError:
                raise
            except Exception:
                pass
            try:
                image_dict: dict = xlsx_embed_cells_images(io.BytesIO(buffer))
                save_image([item for item in image_dict.values()])
            except ResourceLimitError:
                raise
            except Exception as e:
                image_dict = {}
            worksheets = workbook.worksheets
            worksheets_size = len(worksheets)
            return [row for row in
                    [handle_sheet(file.name,
                                  sheet,
                                  image_dict,
                                  limit) if worksheets_size == 1 and sheet.title == 'Sheet1' else handle_sheet(
                        sheet.title, sheet, image_dict, limit) for sheet
                     in worksheets] if row is not None]
        except (SliceError, ResourceLimitError):
            raise
        except Exception as e:
            _log.error(f"Error processing XLSX file {file.name}: {e}, {traceback.format_exc()}")
            # 解析失败显式报错，避免静默返回空段落
            raise SliceError(
                500,
                _("Failed to parse Excel file {name}: {reason}").format(name=file.name, reason=e),
            )

    def get_content(self, file, save_image):
        try:
            # 加载 Excel 文件
            buffer = file.read()
            validate_workbook(buffer, getattr(file, 'name', ''))
            workbook = load_workbook(io.BytesIO(buffer), keep_links=False)
            validate_worksheet_sizes(workbook)
            try:
                image_dict: dict = xlsx_embed_cells_images(io.BytesIO(buffer))
                if len(image_dict) > 0:
                    save_image(image_dict.values())
            except ResourceLimitError:
                raise
            except Exception as e:
                _log.error(f'Exception: {e}')
                image_dict = {}
            md_tables = ''
            # 遍历所有工作表
            for sheetname in workbook.sheetnames:
                sheet = workbook[sheetname]
                rows = self.fill_merged_cells(sheet, image_dict)
                if len(rows) == 0:
                    continue

                # 添加 sheet 名称作为标题
                md_tables += f'## {sheetname}\n\n'

                # 提取表头和内容
                headers = [f"{key}" for key, value in rows[0].items()]

                # 构建 Markdown 表格
                md_table = '| ' + ' | '.join(headers) + ' |\n'
                md_table += '| ' + ' | '.join(['---'] * len(headers)) + ' |\n'
                for row in rows:
                    r = [self._escape_cell_content(value) for key, value in row.items()]
                    md_table += '| ' + ' | '.join(r) + ' |\n'

                md_tables += md_table + '\n\n'

            return md_tables
        except (SliceError, ResourceLimitError):
            raise
        except Exception as e:
            _log.error(f'excel split handle error: {e}')
            raise SliceError(400, 'Invalid or damaged Excel workbook') from e

    def _escape_cell_content(self, cell_value):
        """转义单元格内容,避免破坏 Markdown 表格结构"""
        if cell_value is None:
            return ''

        cell_str = str(cell_value)

        # 替换换行符为 <br>
        cell_str = cell_str.replace('\n', '<br>')

        # 转义管道符 | 为 HTML 实体
        cell_str = cell_str.replace('|', '&#124;')

        # 如果内容包含反引号,需要转义
        if '`' in cell_str:
            cell_str = cell_str.replace('`', '&#96;')

        return cell_str

    def support(self, file, get_buffer):
        file_name: str = file.name.lower()
        if file_name.endswith(tuple(WORKBOOK_CONTENT_TYPES)):
            return True
        return False
