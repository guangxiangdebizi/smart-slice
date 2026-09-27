# coding=utf-8
# Derived from the source platform document-slicing layer (see docs/PORTING.md).
# Imports and framework-facing symbols were re-routed; the parsing and
# splitting logic is unchanged.
from smart_slice._logging import get_logger

_log = get_logger("table.xlsx")


import io
import logging
import traceback

try:
    from openpyxl import load_workbook
except ImportError:  # optional extra: smart-slice[office]
    load_workbook = None

from smart_slice.qa._table_base import BaseParseTableHandle
from smart_slice.handlers._xlsx_images import xlsx_embed_cells_images


class XlsxParseTableHandle(BaseParseTableHandle):
    def support(self, file, get_buffer):
        file_name: str = file.name.lower()
        if file_name.endswith('.xlsx'):
            return True
        return False

    def fill_merged_cells(self, sheet, image_dict):
        data = []

        # 获取第一行作为标题行
        headers = []
        for idx, cell in enumerate(sheet[1]):
            if cell.value is None:
                headers.append(' ' * (idx + 1))
            else:
                headers.append(cell.value)

        # 从第二行开始遍历每一行
        for row in sheet.iter_rows(min_row=2, values_only=False):
            row_data = {}
            for col_idx, cell in enumerate(row):
                cell_value = cell.value

                # 如果单元格为空，并且该单元格在合并单元格内，获取合并单元格的值
                if cell_value is None:
                    for merged_range in sheet.merged_cells.ranges:
                        if cell.coordinate in merged_range:
                            cell_value = sheet[merged_range.min_row][merged_range.min_col - 1].value
                            break
                if cell_value is None:
                    cell_value = ''
                image = image_dict.get(cell_value, None)
                if image is not None:
                    cell_value = f'![](./oss/file/{image.id})'

                # 使用标题作为键，单元格的值作为值存入字典
                row_data[headers[col_idx]] = cell_value
            data.append(row_data)

        return data

    def handle(self, file, get_buffer, save_image):
        buffer = get_buffer(file)
        try:
            wb = load_workbook(io.BytesIO(buffer))
            try:
                image_dict: dict = xlsx_embed_cells_images(io.BytesIO(buffer))
                save_image([item for item in image_dict.values()])
            except Exception as e:
                image_dict = {}
            result = []
            for sheetname in wb.sheetnames:
                paragraphs = []
                ws = wb[sheetname]
                data = self.fill_merged_cells(ws, image_dict)

                for row in data:
                    row_output = "; ".join([f"{key}: {value}" for key, value in row.items()])
                    paragraphs.append({'title': '', 'content': row_output})

                result.append({'name': sheetname, 'paragraphs': paragraphs})

        except BaseException as e:
            _log.error(f"Error processing XLSX file {file.name}: {e}, {traceback.format_exc()}")
            return [{'name': file.name, 'paragraphs': []}]
        return result

    def get_content(self, file, save_image):
        try:
            # 加载 Excel 文件
            workbook = load_workbook(file)
            try:
                image_dict: dict = xlsx_embed_cells_images(file)
                if len(image_dict) > 0:
                    save_image(image_dict.values())
            except Exception as e:
                _log.error(f'Exception: {e}')
                image_dict = {}
            md_tables = ''
            # 遍历所有工作表
            for sheetname in workbook.sheetnames:
                sheet = workbook[sheetname]
                rows = self.fill_merged_cells(sheet, image_dict)
                if len(rows) == 0:
                    continue

                # 添加 sheet 名称作为标题
                md_tables += f'## {sheetname}\n\n'

                # 提取表头和内容
                headers = [f"{key}" for key, value in rows[0].items()]

                # 构建 Markdown 表格
                md_table = '| ' + ' | '.join(headers) + ' |\n'
                md_table += '| ' + ' | '.join(['---'] * len(headers)) + ' |\n'
                for row in rows:
                    r = [f'{value}' for key, value in row.items()]
                    md_table += '| ' + ' | '.join(
                        [str(cell).replace('\n', '<br>') if cell is not None else '' for cell in r]) + ' |\n'

                md_tables += md_table + '\n\n'

            return md_tables
        except Exception as e:
            _log.error(f'excel split handle error: {e}')
            return f'error: {e}'
